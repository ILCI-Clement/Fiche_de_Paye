"""Excel content, authorization and audit checks using disposable synthetic records."""
from io import BytesIO
import json
import unittest
from openpyxl import load_workbook
import test_personnel_register_sensitive as fixture


class PersonnelExcelTests(unittest.TestCase):
    setUp = fixture.PersonnelSensitiveHTTPTests.setUp
    create = fixture.PersonnelSensitiveHTTPTests.create
    rows = fixture.PersonnelSensitiveHTTPTests.rows
    raw_update = fixture.PersonnelSensitiveHTTPTests.raw_update

    def export(self, **extra):
        return self.client.post('/personnel-register/export.xlsx', json={
            'q': '', 'personnel_type': '', 'archive': 'active', 'status': 'all',
            'confirm_sensitive_export': True, **extra,
        })

    def test_all_filtered_pages_and_periods_export_with_literal_text_and_plain_nir(self):
        for index in range(22):
            response = self.create(full_name=f'=MATCH synthetic {index}', social_security_number=fixture.SYNTHETIC_NIR,
                                   birth_date='2000-01-02', initial_period={'entry_date': '2026-01-01', 'contract_type': 'CDI'})
            self.assertEqual(response.status_code, 201, response.text)
        self.create(full_name='Other synthetic person')
        page = self.client.get('/personnel-register?q=MATCH&page_size=20').json()
        self.assertEqual(page['total'], 22)
        self.assertEqual(len(page['employees']), 20)
        response = self.export(q='MATCH')
        self.assertEqual(response.status_code, 200, response.text[:100] if response.status_code != 200 else '')
        self.assertIn('no-store', response.headers['cache-control'])
        book = load_workbook(BytesIO(response.content))
        sheet = book.active
        self.assertEqual(sheet.max_row, 29)
        self.assertEqual(sheet['B8'].data_type, 's')
        self.assertTrue(sheet['B8'].value.startswith('='))
        self.assertEqual(sheet['F8'].value, fixture.SYNTHETIC_NIR)
        self.assertEqual(sheet['F8'].data_type, 's')
        self.assertEqual(sheet['G8'].value.strftime('%Y-%m-%d'), '2000-01-02')
        self.assertEqual(sheet['J8'].value.strftime('%Y-%m-%d'), '2026-01-01')
        book.close()
        events = [row for row in self.rows('personnel_register_events') if row['action'] == 'excel_exported']
        self.assertEqual(len(events), 22)
        self.assertNotIn(fixture.SYNTHETIC_NIR, json.dumps(events))

    def test_permission_confirmation_and_failure_do_not_release_a_file(self):
        self.create(social_security_number=fixture.SYNTHETIC_NIR)
        self.assertEqual(self.export(confirm_sensitive_export=False).status_code, 400)
        self.assertEqual(self.export(confirm_sensitive_export='true').status_code, 422)
        self.assertEqual(self.export(q='no-match').status_code, 404)
        self.actor = {'username': 'fixture-employee', 'role_tags': ['Admin'], 'session_version': 0}
        self.assertEqual(self.export().status_code, 403)
        self.actor = {'username': 'fixture-admin', 'role_tags': ['Admin'], 'session_version': 0}
        self.fail_audit = True
        with self.assertLogs('backend.personnel_register', level='ERROR'):
            response = self.export()
        self.assertEqual(response.status_code, 500)
        self.assertNotIn('attachment', response.headers.get('content-disposition', ''))
        self.assertNotIn(fixture.SYNTHETIC_NIR, response.text)
        self.assertFalse(any(row['action'] == 'excel_exported' for row in self.rows('personnel_register_events')))
