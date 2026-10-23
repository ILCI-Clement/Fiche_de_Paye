"""Employee-register XLSX generation from the supplied blank company template."""
from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


TEMPLATE = Path(__file__).resolve().parents[1] / 'template' / 'Exemple-de-modele-de-registre-unique-du-personnel.xlsx'
MAX_EXPORT_ROWS = 10_000
_CONTRACTS = {'CDI': 'CDI', 'CDD': 'CDD', 'stage': 'Convention de stage'}


def generate_personnel_excel(rows):
    """One row per employment period; never invent unavailable template fields."""
    if not rows or len(rows) > MAX_EXPORT_ROWS:
        raise ValueError('Le nombre de lignes à exporter est invalide.')
    workbook = load_workbook(TEMPLATE, keep_links=False)
    sheet = workbook.active
    styles = [copy(sheet.cell(8, column)._style) for column in range(1, 19)]
    # Remove the template's 150 placeholder numbers; grow from its first styled row.
    sheet.delete_rows(8, max(0, sheet.max_row - 7))
    sheet['E2'] = 'REGISTRE UNIQUE DU PERSONNEL'
    for number, row in enumerate(rows, 1):
        contract = (row.get('other_contract', '') if row.get('contract_type') == 'autre'
                    else _CONTRACTS.get(row.get('contract_type'), ''))
        values = [number, row['full_name'], row.get('nationality'),
                  {'F': 'Femme', 'M': 'Homme'}.get(row.get('sex'), ''), row.get('address'),
                  row.get('social_security_number'), row.get('birth_date'), row.get('job_title'),
                  row.get('qualification'), row.get('entry_date'), row.get('actual_exit_date'),
                  row.get('work_permit_type') if row.get('is_foreign_worker') else '',
                  row.get('work_permit_number') if row.get('is_foreign_worker') else '', contract,
                  '\n'.join(filter(None, [row.get('temporary_company_name'), row.get('temporary_company_address')])) if row.get('is_temporary_worker') else '',
                  'Oui' if row.get('is_young_worker') and row.get('youth_contract_type') == 'apprenti' else '',
                  'Oui' if row.get('is_young_worker') and row.get('youth_contract_type') == 'qualification' else '',
                  'Oui' if row.get('is_young_worker') and row.get('youth_contract_type') == 'adaptation' else '']
        for column, value in enumerate(values, 1):
            cell = sheet.cell(number + 7, column)
            cell._style = copy(styles[column - 1])
            if isinstance(value, date):
                cell.value = value
                cell.number_format = 'dd/mm/yyyy'
            elif isinstance(value, int):
                cell.value = value
            elif value is not None:
                cell.value = ILLEGAL_CHARACTERS_RE.sub('', str(value))
                # Treat employee input as literal text, including leading '='.
                cell.data_type = 's'
                cell.number_format = '@'
        # Long addresses/names remain readable rather than shrinking the template font.
        sheet.row_dimensions[number + 7].height = min(409, max(30,
            15 * max(sum(max(1, (len(line) + width - 1) // width)
                         for line in str(values[index] or '').split('\n'))
                     for index, width in [(1, 30), (4, 40), (7, 24), (8, 24)])))
    # The original column descriptor spans all 16,384 columns; constrain it before sizing.
    sheet.column_dimensions['A'].max = 1
    widths = {'A': 11, 'B': 32, 'C': 18, 'D': 12, 'E': 44, 'F': 23, 'G': 18,
              'H': 26, 'I': 26, 'J': 17, 'K': 17, 'L': 20, 'M': 20, 'N': 25,
              'O': 34, 'P': 17, 'Q': 20, 'R': 20}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'C8'
    sheet.print_title_rows = '1:7'
    sheet.print_title_cols = 'A:B'
    sheet.print_area = f'A1:R{len(rows) + 7}'
    workbook.properties.creator = 'ILCI · Espace RH'
    workbook.properties.lastModifiedBy = 'ILCI · Espace RH'
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
